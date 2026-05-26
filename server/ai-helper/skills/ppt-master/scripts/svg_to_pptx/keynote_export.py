"""Keynote-friendly export: rasterize SVGs (macOS qlmanage) into a picture-only PPTX.

Why:
  Keynote often fails to import complex DrawingML shape decks. A picture-only PPTX is
  reliably importable while keeping high visual fidelity.

Notes:
  - Output slides are NOT editable (they are full-slide PNG pictures).
  - Uses qlmanage, so this requires macOS.
"""

from __future__ import annotations

import re
import subprocess
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu


EMU_PER_PIXEL = 914400 / 96


def _parse_svg_viewbox(svg_path: Path) -> tuple[int, int]:
    text = svg_path.read_text(encoding="utf-8", errors="ignore")[:4000]
    m = re.search(
        r'viewBox\s*=\s*["\']?\s*[\d.]+\s+[\d.]+\s+([\d.]+)\s+([\d.]+)',
        text,
        re.I,
    )
    if m:
        return max(1, int(float(m.group(1)))), max(1, int(float(m.group(2))))
    mw = re.search(r'width\s*=\s*["\']?(\d+)', text, re.I)
    mh = re.search(r'height\s*=\s*["\']?(\d+)', text, re.I)
    if mw and mh:
        return int(mw.group(1)), int(mh.group(1))
    return 1280, 720


def _center_crop_to_aspect(im: Image.Image, target_w: int, target_h: int) -> Image.Image:
    """Crop within bounds to the desired aspect ratio (no padding, no expansion)."""
    iw, ih = im.size
    ratio = target_w / target_h
    if iw / ih > ratio:
        crop_w = int(round(ih * ratio))
        left = max(0, (iw - crop_w) // 2)
        return im.crop((left, 0, left + crop_w, ih))
    crop_h = int(round(iw / ratio))
    top = max(0, (ih - crop_h) // 2)
    return im.crop((0, top, iw, top + crop_h))


def _rasterize_one(
    svg_path: Path,
    out_png: Path,
    *,
    export_w: int,
    export_h: int,
    supersample: int,
    png_compress: int,
) -> None:
    if supersample < 1:
        supersample = 1
    ql_edge = max(export_w, export_h) * supersample
    out_png.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["qlmanage", "-t", f"-s{ql_edge}", "-o", str(out_png.parent), str(svg_path)],
        check=True,
        capture_output=True,
    )
    raw = out_png.parent / f"{svg_path.name}.png"
    if not raw.exists():
        raise FileNotFoundError(f"qlmanage did not produce {raw}")

    im = Image.open(raw)
    im = _center_crop_to_aspect(im, export_w, export_h)
    if im.size != (export_w, export_h):
        im = im.resize((export_w, export_h), Image.Resampling.LANCZOS)
    im.save(out_png, format="PNG", optimize=False, compress_level=png_compress)
    raw.unlink(missing_ok=True)


def export_keynote_pptx(
    svg_files: list[Path],
    output_path: Path,
    *,
    scale: int = 4,
    supersample: int = 2,
    png_compress: int = 6,
) -> None:
    """Export svg_files into a Keynote-friendly picture-only PPTX."""
    if not svg_files:
        raise ValueError("No SVG files provided")
    if scale < 1:
        scale = 1
    if not (0 <= png_compress <= 9):
        raise ValueError("png_compress must be 0..9")

    slide_w, slide_h = _parse_svg_viewbox(svg_files[0])
    export_w, export_h = slide_w * scale, slide_h * scale

    prs = Presentation()
    prs.slide_width = Emu(int(slide_w * EMU_PER_PIXEL))
    prs.slide_height = Emu(int(slide_h * EMU_PER_PIXEL))
    blank = prs.slide_layouts[6]

    with tempfile.TemporaryDirectory(prefix="ppt_keynote_") as tmp:
        tmp_dir = Path(tmp)
        for svg in svg_files:
            png = tmp_dir / f"{svg.stem}.png"
            _rasterize_one(
                svg,
                png,
                export_w=export_w,
                export_h=export_h,
                supersample=supersample,
                png_compress=png_compress,
            )
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(
                str(png), 0, 0, width=prs.slide_width, height=prs.slide_height
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))

